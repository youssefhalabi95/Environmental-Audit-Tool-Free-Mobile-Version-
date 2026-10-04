from flask import Flask, render_template, request, send_file
from openpyxl import Workbook
from io import BytesIO
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import cm

app = Flask(__name__)


CHECKLIST = {

    "Waste Management": [
        "Are waste streams properly identified and segregated?",
        "Are waste containers clearly labelled according to the waste type?",
        "Are waste storage areas appropriately designated and maintained?",
        "Are hazardous and non-hazardous wastes stored separately?",
        "Are waste storage areas protected from rain, runoff and unauthorized access?",
        "Are waste collection and disposal activities properly documented?",
        "Are wastes transferred only to authorized waste contractors or facilities?",
        "Are waste manifests, transfer notes or disposal records maintained?",
        "Are waste containers in good condition and free from significant damage or leakage?",
        "Are opportunities for waste reduction, reuse and recycling being implemented?"
    ],

    "Hazardous Materials": [
        "Are hazardous materials stored in designated areas?",
        "Are chemical containers properly labelled?",
        "Are Safety Data Sheets (SDS) available for applicable chemicals?",
        "Are incompatible chemicals stored separately?",
        "Are secondary containment measures provided where required?",
        "Are chemical storage areas protected from unauthorized access?",
        "Are spill response materials available near chemical storage areas?",
        "Are chemical containers inspected regularly for leakage or damage?",
        "Are expired or unused chemicals properly managed and disposed of?",
        "Are employees handling hazardous materials appropriately trained?"
    ],

    "Spill Prevention & Response": [
        "Are spill prevention measures implemented for fuels, oils and chemicals?",
        "Are spill kits available at relevant locations?",
        "Are spill kits adequately stocked and accessible?",
        "Are employees aware of the spill response procedure?",
        "Are significant spills and environmental incidents documented?",
        "Are contaminated materials generated from spills properly managed?",
        "Are drainage systems protected from potential chemical or oil releases?",
        "Are fuel and chemical transfer activities carried out in controlled areas?"
    ],

    "Air Quality & Dust": [
        "Are appropriate measures implemented to control dust generation?",
        "Are roads and exposed surfaces maintained to minimize dust?",
        "Are dusty materials appropriately covered during transportation?",
        "Are material stockpiles appropriately managed?",
        "Are vehicle speeds controlled in areas where dust generation is significant?",
        "Are air emission sources identified and appropriately controlled?",
        "Are required air quality or emission monitoring activities being conducted?",
        "Are monitoring results reviewed and corrective actions implemented where necessary?"
    ],

    "Noise": [
        "Are significant noise sources identified?",
        "Are noise-generating activities appropriately controlled?",
        "Are noisy activities restricted to appropriate working hours where required?",
        "Are equipment and machinery maintained to minimize excessive noise?",
        "Are noise monitoring activities conducted where required?",
        "Are noise complaints recorded and investigated?",
        "Are appropriate mitigation measures implemented where monitoring indicates elevated noise?"
    ],

    "Water & Wastewater": [
        "Are water consumption sources identified and monitored?",
        "Are wastewater discharge points identified?",
        "Are wastewater discharges managed in accordance with applicable requirements?",
        "Are wastewater treatment systems properly operated and maintained?",
        "Are wastewater monitoring results documented?",
        "Are stormwater drainage systems maintained and free from obstruction?",
        "Are potentially contaminated runoff pathways appropriately controlled?",
        "Are leaks or unauthorized discharges promptly investigated and addressed?"
    ],

    "Soil & Groundwater": [
        "Are activities with potential soil contamination risks identified?",
        "Are fuel and chemical storage areas appropriately protected?",
        "Are spills and leaks promptly cleaned and investigated?",
        "Are contaminated soils properly identified and managed?",
        "Are excavation and soil handling activities appropriately controlled?",
        "Are areas with historical contamination appropriately monitored or managed?"
    ],

    "Biodiversity": [
        "Have sensitive ecological receptors in the project area been identified?",
        "Are activities restricted in environmentally sensitive areas where required?",
        "Are vegetation clearance activities controlled and authorized?",
        "Are measures implemented to minimize impacts on wildlife?",
        "Are biodiversity-related monitoring requirements being implemented?",
        "Are ecological incidents or wildlife interactions documented?"
    ],

    "Environmental Permits & Legal Compliance": [
        "Are all required environmental permits and approvals available?",
        "Are permits and approvals valid and up to date?",
        "Are permit conditions being monitored for compliance?",
        "Are environmental regulatory reporting requirements being fulfilled?",
        "Are required environmental records maintained?",
        "Are applicable environmental laws and regulations identified?",
        "Are changes to applicable environmental requirements monitored?",
        "Are regulatory inspections and correspondence properly documented?"
    ],

    "Environmental Management System": [
        "Is there a documented environmental policy?",
        "Are environmental responsibilities clearly assigned?",
        "Are environmental objectives and targets established where appropriate?",
        "Are environmental procedures available to relevant personnel?",
        "Are environmental training records maintained?",
        "Are environmental incidents documented and investigated?",
        "Are corrective and preventive actions tracked to completion?",
        "Are environmental records properly maintained and controlled?",
        "Are environmental performance results periodically reviewed?"
    ],

    "Emergency Preparedness": [
        "Are environmental emergency scenarios identified?",
        "Is an environmental emergency response procedure available?",
        "Are emergency contact details readily available?",
        "Are spill response materials accessible?",
        "Are relevant employees trained in environmental emergency response?",
        "Are emergency drills conducted where required?",
        "Are lessons learned from environmental incidents incorporated into procedures?"
    ]
}


@app.route("/")
def landing():
    return render_template(
        "landing.html",
        categories=list(CHECKLIST.keys())
    )


@app.route("/tool")
def tool():
    return render_template(
        "index.html",
        categories=list(CHECKLIST.keys()),
        checklist=CHECKLIST
    )


def get_audit_data():

    questions = request.form.getlist("question")
    categories = request.form.getlist("category")
    statuses = request.form.getlist("status")
    findings = request.form.getlist("finding")
    actions = request.form.getlist("action")
    responsible = request.form.getlist("responsible")
    target_dates = request.form.getlist("target_date")

    rows = []

    for i, question in enumerate(questions):

        rows.append({
            "no": i + 1,
            "category": categories[i] if i < len(categories) else "",
            "question": question,
            "status": statuses[i] if i < len(statuses) else "",
            "finding": findings[i] if i < len(findings) else "",
            "action": actions[i] if i < len(actions) else "",
            "responsible": responsible[i] if i < len(responsible) else "",
            "target_date": target_dates[i] if i < len(target_dates) else ""
        })

    return rows


def get_summary(rows):

    compliant = sum(
        r["status"] == "Compliant"
        for r in rows
    )

    partial = sum(
        r["status"] == "Partially Compliant"
        for r in rows
    )

    non_compliant = sum(
        r["status"] == "Non-Compliant"
        for r in rows
    )

    na = sum(
        r["status"] == "N/A"
        for r in rows
    )

    applicable = compliant + partial + non_compliant

    if applicable:
        score = round(
            ((compliant + partial * 0.5) / applicable) * 100,
            1
        )
    else:
        score = 0

    return compliant, partial, non_compliant, na, score


# ---------------------------------------------------------
# EXCEL EXPORT
# ---------------------------------------------------------

@app.route("/export/excel", methods=["POST"])
def export_excel():

    rows = get_audit_data()

    project_name = request.form.get(
        "project_name", ""
    )

    site_name = request.form.get(
        "site_name", ""
    )

    auditor_name = request.form.get(
        "auditor_name", ""
    )

    audit_date = request.form.get(
        "audit_date", ""
    )

    selected_categories = request.form.getlist(
        "selected_category"
    )

    wb = Workbook()

    ws = wb.active

    ws.title = "Audit Summary"

    ws["A1"] = "Environmental Audit Checklist"

    info = [

        ("Project", project_name),

        ("Site / Location", site_name),

        ("Auditor", auditor_name),

        ("Audit Date", audit_date),

        (
            "Selected Categories",
            ", ".join(selected_categories)
        )

    ]

    for i, (label, value) in enumerate(
        info,
        start=3
    ):

        ws.cell(i, 1).value = label
        ws.cell(i, 2).value = value


    compliant, partial, non_compliant, na, score = \
        get_summary(rows)


    summary = [

        ("Compliant", compliant),

        ("Partially Compliant", partial),

        ("Non-Compliant", non_compliant),

        ("N/A", na),

        ("Compliance Score", f"{score}%")

    ]


    for i, (label, value) in enumerate(
        summary,
        start=9
    ):

        ws.cell(i, 1).value = label
        ws.cell(i, 2).value = value


    detail = wb.create_sheet(
        "Audit Findings"
    )


    headers = [

        "No.",

        "Category",

        "Audit Question",

        "Status",

        "Finding / Observation",

        "Corrective Action",

        "Responsible Person",

        "Target Date"

    ]


    for col, header in enumerate(
        headers,
        start=1
    ):

        detail.cell(
            1,
            col
        ).value = header


    for row_num, item in enumerate(
        rows,
        start=2
    ):

        values = [

            item["no"],

            item["category"],

            item["question"],

            item["status"],

            item["finding"],

            item["action"],

            item["responsible"],

            item["target_date"]

        ]


        for col, value in enumerate(
            values,
            start=1
        ):

            detail.cell(
                row_num,
                col
            ).value = value


    for sheet in wb.worksheets:

        for column in sheet.columns:

            max_length = max(
                (
                    len(str(cell.value))
                    for cell in column
                    if cell.value
                ),
                default=0
            )

            sheet.column_dimensions[
                column[0].column_letter
            ].width = min(
                max_length + 2,
                50
            )

        sheet.freeze_panes = "A2"


    output = BytesIO()

    wb.save(output)

    output.seek(0)


    return send_file(

        output,

        as_attachment=True,

        download_name=
        "Environmental_Audit_Checklist.xlsx",

        mimetype=
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    )


# ---------------------------------------------------------
# PDF EXPORT
# ---------------------------------------------------------

@app.route("/export/pdf", methods=["POST"])
def export_pdf():

    rows = get_audit_data()

    project_name = request.form.get(
        "project_name", ""
    )

    site_name = request.form.get(
        "site_name", ""
    )

    auditor_name = request.form.get(
        "auditor_name", ""
    )

    audit_date = request.form.get(
        "audit_date", ""
    )

    selected_categories = request.form.getlist(
        "selected_category"
    )


    compliant, partial, non_compliant, na, score = \
        get_summary(rows)


    output = BytesIO()


    doc = SimpleDocTemplate(

        output,

        pagesize=A4,

        rightMargin=1.3 * cm,

        leftMargin=1.3 * cm,

        topMargin=1.3 * cm,

        bottomMargin=1.3 * cm

    )


    styles = getSampleStyleSheet()

    title = styles["Title"]

    heading = styles["Heading2"]

    body = styles["BodyText"]

    body.leading = 13


    story = [

        Paragraph(
            "Environmental Audit Checklist",
            title
        ),

        Spacer(1, 10),

        Paragraph(
            f"<b>Project:</b> {project_name}",
            body
        ),

        Paragraph(
            f"<b>Site / Location:</b> {site_name}",
            body
        ),

        Paragraph(
            f"<b>Auditor:</b> {auditor_name}",
            body
        ),

        Paragraph(
            f"<b>Audit Date:</b> {audit_date}",
            body
        ),

        Paragraph(
            f"<b>Categories:</b> "
            f"{', '.join(selected_categories)}",
            body
        ),

        Spacer(1, 10),

        Paragraph(

            f"<b>Summary:</b> "
            f"Compliant: {compliant} | "
            f"Partial: {partial} | "
            f"Non-Compliant: {non_compliant} | "
            f"N/A: {na} | "
            f"Score: {score}%",

            body

        ),

        Spacer(1, 15)

    ]


    current_category = None


    for item in rows:

        if item["category"] != current_category:

            current_category = item["category"]

            story.append(
                Paragraph(
                    current_category,
                    heading
                )
            )

            story.append(
                Spacer(1, 5)
            )


        data = [

            [

                Paragraph(
                    "<b>Audit Question</b>",
                    body
                ),

                Paragraph(
                    item["question"],
                    body
                )

            ],

            [

                Paragraph(
                    "<b>Status</b>",
                    body
                ),

                Paragraph(
                    item["status"] or "Not assessed",
                    body
                )

            ],

            [

                Paragraph(
                    "<b>Finding</b>",
                    body
                ),

                Paragraph(
                    item["finding"] or "-",
                    body
                )

            ],

            [

                Paragraph(
                    "<b>Corrective Action</b>",
                    body
                ),

                Paragraph(
                    item["action"] or "-",
                    body
                )

            ],

            [

                Paragraph(
                    "<b>Responsible</b>",
                    body
                ),

                Paragraph(
                    item["responsible"] or "-",
                    body
                )

            ],

            [

                Paragraph(
                    "<b>Target Date</b>",
                    body
                ),

                Paragraph(
                    item["target_date"] or "-",
                    body
                )

            ]

        ]


        table = Table(

            data,

            colWidths=[
                3.3 * cm,
                14.7 * cm
            ]

        )


        table.setStyle(

            TableStyle([

                (
                    "GRID",
                    (0, 0),
                    (-1, -1),
                    0.5,
                    colors.grey
                ),

                (
                    "VALIGN",
                    (0, 0),
                    (-1, -1),
                    "TOP"
                ),

                (
                    "BACKGROUND",
                    (0, 0),
                    (0, -1),
                    colors.whitesmoke
                ),

                (
                    "LEFTPADDING",
                    (0, 0),
                    (-1, -1),
                    5
                ),

                (
                    "RIGHTPADDING",
                    (0, 0),
                    (-1, -1),
                    5
                ),

                (
                    "TOPPADDING",
                    (0, 0),
                    (-1, -1),
                    5
                ),

                (
                    "BOTTOMPADDING",
                    (0, 0),
                    (-1, -1),
                    5
                )

            ])

        )


        story.append(table)

        story.append(
            Spacer(1, 10)
        )


    doc.build(story)

    output.seek(0)


    return send_file(

        output,

        as_attachment=True,

        download_name=
        "Environmental_Audit_Checklist.pdf",

        mimetype="application/pdf"

    )


# ---------------------------------------------------------
# ABOUT
# ---------------------------------------------------------

@app.route("/about")
def about():

    return render_template(
        "about.html"
    )


# ---------------------------------------------------------
# PRIVACY
# ---------------------------------------------------------

@app.route("/privacy")
def privacy():

    return render_template(
        "privacy.html"
    )


if __name__ == "__main__":

    app.run(
        debug=True
    )